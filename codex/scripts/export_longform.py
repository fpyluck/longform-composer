#!/usr/bin/env python3
"""Export a longform-composer final Markdown file to readable document formats.

The exporter is a low-cost post-processing step. It converts the canonical
merged Markdown without summarizing or rewriting content.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Iterable, List

from longform_common import load_project, write_progress_log

DEFAULT_FORMATS = ["docx"]
PANDOC_FORMATS = {"docx", "html", "pdf"}


def default_source(root: Path) -> Path:
    manifest = load_project(root)
    artifacts = manifest.get("artifacts") if isinstance(manifest.get("artifacts"), dict) else {}
    final = str(artifacts.get("final") or "final/final_merged.md")
    return root / final


def output_path(output_dir: Path, source: Path, fmt: str) -> Path:
    return output_dir / f"{source.stem}.{fmt}"


def run_pandoc(source: Path, output: Path) -> None:
    cmd = [
        "pandoc",
        str(source),
        "--from",
        "gfm",
        "--standalone",
        "--wrap=none",
        "--resource-path",
        str(source.parent),
        "-o",
        str(output),
    ]
    subprocess.run(cmd, check=True)


def markdown_blocks(text: str) -> Iterable[tuple[str, str | tuple[str, List[List[str]]]]]:
    lines = text.splitlines()
    idx = 0
    while idx < len(lines):
        line = lines[idx]
        if not line.strip():
            idx += 1
            continue
        if line.startswith("```"):
            fence: List[str] = []
            idx += 1
            while idx < len(lines) and not lines[idx].startswith("```"):
                fence.append(lines[idx])
                idx += 1
            idx += 1
            yield "code", "\n".join(fence)
            continue
        if re.match(r"^#{1,6}\s+", line):
            yield "heading", line
            idx += 1
            continue
        if re.match(r"^\s*[-*+]\s+", line):
            yield "bullet", re.sub(r"^\s*[-*+]\s+", "", line).strip()
            idx += 1
            continue
        if re.match(r"^\s*\d+[.)]\s+", line):
            yield "number", re.sub(r"^\s*\d+[.)]\s+", "", line).strip()
            idx += 1
            continue
        if line.lstrip().startswith(">"):
            quote_lines: List[str] = []
            while idx < len(lines) and lines[idx].lstrip().startswith(">"):
                quote_lines.append(lines[idx].lstrip()[1:].strip())
                idx += 1
            yield "quote", " ".join(quote_lines).strip()
            continue
        if "|" in line and idx + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[idx + 1]):
            header = [cell.strip() for cell in line.strip().strip("|").split("|")]
            rows: List[List[str]] = [header]
            idx += 2
            while idx < len(lines) and "|" in lines[idx] and lines[idx].strip():
                rows.append([cell.strip() for cell in lines[idx].strip().strip("|").split("|")])
                idx += 1
            yield "table", ("", rows)
            continue

        para: List[str] = [line.strip()]
        idx += 1
        while idx < len(lines) and lines[idx].strip() and not re.match(r"^#{1,6}\s+", lines[idx]) and not lines[idx].startswith("```"):
            if re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[idx]) or lines[idx].lstrip().startswith(">"):
                break
            para.append(lines[idx].strip())
            idx += 1
        yield "paragraph", " ".join(para)


def strip_inline_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("*", "").replace("_", "")
    return text.strip()


def run_python_docx(source: Path, output: Path) -> None:
    try:
        from docx import Document  # type: ignore
        from docx.shared import Pt  # type: ignore
    except ImportError as exc:
        raise RuntimeError("python-docx is required when pandoc is unavailable") from exc

    document = Document()
    for kind, payload in markdown_blocks(source.read_text(encoding="utf-8")):
        if kind == "heading":
            text = str(payload)
            hashes, heading = text.split(" ", 1)
            document.add_heading(strip_inline_markdown(heading), level=min(len(hashes), 4))
        elif kind == "bullet":
            document.add_paragraph(strip_inline_markdown(str(payload)), style="List Bullet")
        elif kind == "number":
            document.add_paragraph(strip_inline_markdown(str(payload)), style="List Number")
        elif kind == "quote":
            paragraph = document.add_paragraph(strip_inline_markdown(str(payload)))
            paragraph.style = "Intense Quote"
        elif kind == "code":
            paragraph = document.add_paragraph(str(payload))
            for run in paragraph.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        elif kind == "table":
            _, rows = payload  # type: ignore[misc]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx in range(width):
                    table.cell(r_idx, c_idx).text = strip_inline_markdown(row[c_idx]) if c_idx < len(row) else ""
        else:
            document.add_paragraph(strip_inline_markdown(str(payload)))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def export_one(root: Path, source: Path, output_dir: Path, fmt: str, engine: str) -> Path:
    fmt = fmt.lower().lstrip(".")
    output = output_path(output_dir, source, fmt)
    output.parent.mkdir(parents=True, exist_ok=True)

    if engine in {"auto", "pandoc"} and fmt in PANDOC_FORMATS and shutil.which("pandoc"):
        run_pandoc(source, output)
        return output
    if fmt == "docx" and engine in {"auto", "python-docx"}:
        run_python_docx(source, output)
        return output

    if engine == "pandoc" and not shutil.which("pandoc"):
        raise RuntimeError("pandoc was requested but is not available on PATH")
    raise RuntimeError(f"unsupported export format or engine: format={fmt}, engine={engine}")


def export(root: Path, formats: List[str], source: Path | None, output_dir: Path | None, engine: str) -> List[Path]:
    root = root.resolve()
    source = source.resolve() if source else default_source(root).resolve()
    if not source.exists():
        raise FileNotFoundError(f"source Markdown not found: {source}")
    output_dir = (output_dir or root / "exports").resolve()

    outputs: List[Path] = []
    for fmt in formats:
        outputs.append(export_one(root, source, output_dir, fmt, engine))
    write_progress_log(root, "exported " + ", ".join(path.name for path in outputs))
    return outputs


def main() -> int:
    parser = argparse.ArgumentParser(description="export longform final Markdown to document formats")
    parser.add_argument("--root", required=True, help="project root containing manifest.yaml")
    parser.add_argument("--source", help="source Markdown path; defaults to manifest artifacts.final")
    parser.add_argument("--output-dir", help="output directory; defaults to <root>/exports")
    parser.add_argument("--format", action="append", dest="formats", choices=["docx", "html", "pdf"], help="output format; repeatable")
    parser.add_argument("--engine", choices=["auto", "pandoc", "python-docx"], default="auto", help="conversion engine")
    args = parser.parse_args()

    outputs = export(
        Path(args.root),
        args.formats or DEFAULT_FORMATS,
        Path(args.source) if args.source else None,
        Path(args.output_dir) if args.output_dir else None,
        args.engine,
    )
    print("Longform export complete")
    for path in outputs:
        print(f"  output: {path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1)
